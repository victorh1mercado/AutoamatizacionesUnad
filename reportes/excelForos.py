from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook


# ---------------------------------------------------------
# Ruta del reporte
# ---------------------------------------------------------

CARPETA_REPORTES = Path("reportes")
CARPETA_REPORTES.mkdir(exist_ok=True)

ARCHIVO_REPORTE = CARPETA_REPORTES / "Reporte_Foros.xlsx"


# ---------------------------------------------------------
# Crear reporte si no existe
# ---------------------------------------------------------

def crear_reporte():

    if ARCHIVO_REPORTE.exists():
        return

    wb = Workbook()
    ws = wb.active

    ws.title = "Publicaciones"

    ws.append([
        "Fecha",
        "Grupo",
        "Discusión",
        "Estado",
        "Observación",
        "URL"
    ])

    wb.save(ARCHIVO_REPORTE)

    print("📄 Reporte creado correctamente.")


# ---------------------------------------------------------
# Registrar un grupo
# ---------------------------------------------------------

def registrar_grupo(
    grupo,
    discusion,
    estado,
    observacion="",
    url=""
):

    crear_reporte()

    wb = load_workbook(ARCHIVO_REPORTE)

    ws = wb["Publicaciones"]

    ws.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        grupo,
        discusion,
        estado,
        observacion,
        url
    ])

    wb.save(ARCHIVO_REPORTE)


# ---------------------------------------------------------
# Resumen final
# ---------------------------------------------------------

def imprimir_resumen():

    wb = load_workbook(ARCHIVO_REPORTE)

    ws = wb["Publicaciones"]

    total = ws.max_row - 1

    print("\n====================================")
    print(" REPORTE GENERADO ")
    print("====================================")
    print(f"Registros : {total}")
    print(f"Archivo   : {ARCHIVO_REPORTE}")
    print("====================================")