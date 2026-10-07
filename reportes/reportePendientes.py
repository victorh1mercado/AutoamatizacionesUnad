from openpyxl import Workbook, load_workbook
from pathlib import Path

RUTA = Path("reportes/Reporte_Pendientes.xlsx")


def crear_reporte_pendientes():

    if RUTA.exists():
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Pendientes"

    ws.append([
        "Grupo",
        "Estudiante",
        "Fecha publicación",
        "Horas transcurridas",
        "Estado",
        "Debe responder",
        "URL"
    ])

    wb.save(RUTA)


def registrar_pendiente(
    grupo,
    estudiante,
    fecha,
    horas,
    estado,
    responder,
    url
):

    wb = load_workbook(RUTA)
    ws = wb.active

    ws.append([
        grupo,
        estudiante,
        fecha,
        round(horas, 2),
        estado,
        responder,
        url
    ])

    wb.save(RUTA)